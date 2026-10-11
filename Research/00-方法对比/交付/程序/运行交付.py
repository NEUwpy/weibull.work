"""Research00 当前交付入口。默认核对，生成操作须显式选择。"""
from pathlib import Path
import argparse
import csv
import hashlib
import importlib.util
import json
import math
import os
import re
import subprocess
import sys
import tempfile

from 共用 import ROOT, PROGRAM, RESULT, matrix_path, setup_runtime

setup_runtime()
CATALOG = PROGRAM / '跨参数对比/13组合取数.json'


def load_json(path):
    return json.loads(path.read_text(encoding='utf-8-sig'))


def module(path, name):
    sys.path.insert(0, str(path.parent))
    spec = importlib.util.spec_from_file_location(name, path)
    obj = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(obj)
    return obj


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def sources(group=None):
    rows = load_json(CATALOG)['sources']
    selected = [s for s in rows if group is None or s['program'] == group]
    if not selected:
        raise ValueError(f'未知参数组：{group}')
    return selected


def verify(group=None):
    from openpyxl import load_workbook
    count = 0
    for source in sources(group):
        unit = PROGRAM / source['program']
        path = matrix_path(source)
        assert digest(path) == source['matrix_sha256'], path
        data = load_json(path)
        assert len(data['samples']) == 150 and len(data['results']) == 1200
        estimates = {(r['n'], r['id'], r['method_id'], r['delta']): r for r in data['results']}
        samples = {(s['n'], s['id']): s['values'] for s in data['samples']}
        assert len(estimates) == 1200 and len(samples) == 150
        old = (unit / '完整矩阵').exists()
        cfg = load_json(unit / ('完整矩阵/config.json' if old else 'config.json'))
        for rel, expected in cfg['source_files'].items():
            assert digest(unit / '输入原件' / rel) == expected, (unit.name, rel)
        for rel, expected in cfg['dependency_hashes'].items():
            assert digest(unit / '依赖快照/python' / rel) == expected, (unit.name, rel)
        book = RESULT / unit.name / (unit.name + '.xlsx')
        wb = load_workbook(book, read_only=True, data_only=True)
        order = [('mdm', .1), ('mdm', .15), ('mdm', .2), ('lse', None),
                 ('lre', None), ('wmle', None), ('mle', None), ('mmle', None)]
        assert wb.sheetnames == [f'{prefix}_n{n}' for n in (7, 15, 30)
                                 for prefix in ('生成样本', '估计结果')]
        for n in (7, 15, 30):
            a = list(wb[f'生成样本_n{n}'].values)
            b = list(wb[f'估计结果_n{n}'].values)
            assert len(a) == 51 and len(b) == 52
            assert [r[0] for r in a[1:]] == [r[0] for r in b[2:]] == list(range(1, 51))
            for sid in range(1, 51):
                assert list(a[sid][1:]) == samples[n, sid]
                for k, (method, delta) in enumerate(order):
                    r = estimates[n, sid, method, delta]
                    vals = [r[v] for v in ('beta_hat', 'eta_hat', 'gamma_hat')]
                    finite = bool(r['converged'] and all(isinstance(v, (int, float))
                                                         and math.isfinite(v) for v in vals))
                    assert list(b[sid + 1][1 + 3*k:4 + 3*k]) == (vals if finite else ['—']*3), (unit.name, n, sid, method)
        wb.close()
        assert len(list((RESULT / unit.name).glob('*.png'))) == 11
        count += 1
    for report in ROOT.glob('报告*.md'):
        for link in re.findall(r'\]\(([^)]+)\)', report.read_text(encoding='utf-8')):
            if not link.startswith(('https:', 'http:', '#')):
                assert (ROOT / link.split('#')[0]).exists(), (report.name, link)
    print(f'PASS：{count}组；每组150样本、1200估计、6张工作表和11张图；输入、方法快照及报告链接通过核对。')


def python_call(path):
    env = dict(os.environ, PYTHONPATH=str(PROGRAM) + os.pathsep + os.environ.get('PYTHONPATH', ''),
               PYTHONUTF8='1', PYTHONDONTWRITEBYTECODE='1')
    # 给子进程补充本机已有绘图库，估计仍由项目 Python 运行。
    env['PYTHONPATH'] += os.pathsep + os.pathsep.join(p for p in sys.path if 'site-packages' in p)
    subprocess.run([sys.executable, '-B', str(path)], env=env, check=True)


def build_workbook(script):
    """运行原建表程序，共用本机的 Node 依赖，避免复制 node_modules。"""
    runtime = Path('C:/Users/36089/.cache/codex-runtimes/codex-primary-runtime/dependencies/node')
    package = runtime / 'node_modules/@oai/artifact-tool'
    entry = package / load_json(package / 'package.json')['exports']['.']
    text = script.read_text(encoding='utf-8-sig')
    text = text.replace("'@oai/artifact-tool'", json.dumps(entry.resolve().as_uri()))
    text = text.replace('path.dirname(fileURLToPath(import.meta.url))', json.dumps(str(script.parent)))
    # 临时入口仅解析依赖与目录；工作簿生成代码保持原有计算和样式。
    with tempfile.TemporaryDirectory(prefix='r00-workbook-') as temporary:
        launcher = Path(temporary) / 'build.mjs'
        launcher.write_text(text, encoding='utf-8')
        subprocess.run([str(runtime / 'bin/node.exe'), str(launcher)], check=True)


def generate(group=None):
    """从冻结输入和缓存重建；未缓存的计算由原批次 runner 执行。"""
    for source in sources(group):
        unit = PROGRAM / source['program']
        out = RESULT / unit.name
        out.mkdir(parents=True, exist_ok=True)
        if (unit / '完整矩阵').exists():
            python_call(unit / '完整矩阵/run_20261007.py')
            build_workbook(unit / '完整矩阵/build_matrix.mjs')
            python_call(unit / '分布与统计/read_excel.py')
            python_call(unit / '分布与统计/draw.py')
        else:
            python_call(unit / '本次补算.py')
            build_workbook(unit / '本次建表.mjs')
            python_call(unit / '核对本次表.py')
            python_call(unit / '本次分布图.py')
            python_call(unit / '本次MDM图.py')
    verify(group)


def report1():
    folder = PROGRAM / '报告1'
    module(folder / '计算逐法过程曲线.py', 'r1_profiles').main()
    module(folder / '推导判据与统一单例.py', 'r1_examples').main()
    module(folder / '构建公式与拟合点数据.py', 'r1_formula_points').main()
    module(folder / '绘制图1小提琴.py', 'r1_violin').main()
    intro = module(folder / '20261010-导入样本起点图.py', 'r1_intro')
    cached = load_json(folder / '20261010-导入图n30样本.json')
    intro.main(extra_samples=cached['samples'], extra_provenance=cached.get('provenance'),
               minimum_distributions='theoretical')
    module(folder / '绘制样本相容分布.py', 'r1_compatibility').main()
    charts = module(folder / '绘制公式与拟合点图.py', 'r1_methods')
    charts.main()
    module(folder / '绘制寿命对照图.py', 'r1_life').main()
    extra = module(folder / '绘制补充图.py', 'r1_supplement')
    extra.continuous_estimates()
    extra.compensation()
    extra.supplement()
    print('报告1：当前主图及补充图已由保存的样本、估计和过程数据生成。')


def report2():
    folder = PROGRAM / '报告2'
    data = RESULT / '报告2/数据'
    diagnostics = module(folder / '逐方法机制复核.py', 'r2_diagnostics')
    diagnostics.refresh_saved_diagnostics()
    diagnostics.audit_mle_case37()
    module(folder / '第一章频率.py', 'r2_frequency').main()
    heatmaps = module(folder / '20261010-逐方法异常分类三子图.py', 'r2_classes')
    qa = load_json(data / '20261010-异常分类三档热力表核验.json')
    counts = {(r['method'], r['case'], r['combination'], int(r['n'])): int(r['count'])
              for r in csv.DictReader((data / '20261010-各类异常参数空间.csv').open(encoding='utf-8-sig'))}
    assert len(counts) == 624 and sum(counts.values()) == 963
    for method in heatmaps.METHODS:
        cases = [r for r in qa['row_order'] if r['method'] == method]
        combinations = qa['column_order']
        matrices = {str(n): __import__('numpy').array([[counts[method, c['case'], combo, n]
                                                      for c in cases] for combo in combinations])
                    for n in heatmaps.NS}
        heatmaps.render(method, matrices, cases, combinations)
    module(folder / '20261010-配对图参照线精简.py', 'r2_pairs').main()
    pairs = load_json(data / '20261010-配对图参照线精简核验.json')
    rows = list(csv.DictReader((data / '20261010-参照线精简图数据.csv').open(encoding='utf-8-sig')))
    lre = module(folder / '20261010-LRE相关系数判据.py', 'r2_lre')
    style = lre.load_style()
    for label in lre.LABELS:
        lre.render(label, pairs['pairs'][label], rows, pairs['figures'][label], style)
    module(folder / '20261010-MDM条件尺度图.py', 'r2_mdm_eta').main()
    module(folder / '20261010-异常分流与条件形状.py', 'r2_fixed').fixed_figure()
    print('报告2：24张当前图已由13组矩阵及保存的过程数据生成。')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('action', nargs='?', default='verify',
                        choices=['verify', 'generate', 'report1', 'report2', 'cross'])
    parser.add_argument('--group', help='例如 W2,1000,3000；默认全部13组')
    args = parser.parse_args()
    if args.action == 'verify':
        verify(args.group)
    elif args.action == 'generate':
        generate(args.group)
    elif args.action == 'report1':
        report1()
    elif args.action == 'report2':
        report2()
    else:
        module(PROGRAM / '跨参数对比/绘制跨参数对比.py', 'r00_cross').main()
