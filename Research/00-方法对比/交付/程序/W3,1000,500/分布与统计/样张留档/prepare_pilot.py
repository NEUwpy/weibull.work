"""Read pilot Excel under the confirmed display rule; check matrix without modifying it."""
import argparse, hashlib, json, math, struct
from pathlib import Path
import openpyxl

def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def same_number(a, b):
    return isinstance(a, (int, float)) and not isinstance(a, bool) and math.isfinite(a) and \
        isinstance(b, (int, float)) and not isinstance(b, bool) and math.isfinite(b) and float(a) == float(b)

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--program', type=Path, default=Path(__file__).resolve().parent)
    args = parser.parse_args(); out = args.program.resolve()
    cfg = json.loads((out/'config.json').read_text(encoding='utf-8'))
    batch = out.parent.parent
    workbook = (batch/cfg['workbook']).resolve()
    matrix_path = (batch/cfg['crosscheck_matrix']).resolve()
    matrix = json.loads(matrix_path.read_text(encoding='utf-8'))
    indexed = {(r['n'], r['id'], r['method_id'], r['delta']): r for r in matrix['results']}
    samples_matrix = {(s['n'], s['id']): s['values'] for s in matrix['samples']}
    assert len(indexed) == 1200 and len(samples_matrix) == 150
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    errors = []; acknowledged = []; records = []; samples = []; counts = []; compared = 0; sample_cells = 0
    approved = cfg['approved_matrix_status_difference']
    approved_key = (approved['n'], approved['id'], approved['method_id'], approved['delta'])
    assert cfg['success_rule'] == 'excel_converged_display'
    for n in cfg['sample_sizes']:
        ws = wb[f'估计结果_n{n}']; rows = list(ws.iter_rows(values_only=True))
        assert rows[0][0] == '样本' and len(rows) == 52
        assert list(rows[1][1:]) == ['α', 'β', 'γ']*8
        headers = [rows[0][1+3*i] for i in range(8)]
        assert len(set(headers)) == 8 and set(headers) == {m['excel_label'] for m in cfg['methods']}
        assert [row[0] for row in rows[2:]] == list(range(1, 51))
        for row in rows[2:]:
            id_ = int(row[0])
            for method in cfg['methods']:
                i = headers.index(method['excel_label'])
                values = list(row[1+3*i:4+3*i])
                success = all(isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v) for v in values)
                failure = all(v in (None, '—', '–', '-') for v in values)
                key = (n, id_, method['method_id'], method['delta']); ref = indexed[key]
                if not success and not failure:
                    errors.append(dict(key=key, issue='partial or unknown Excel triplet', values=values))
                if success != bool(ref['converged']):
                    errors.append(dict(key=key, issue='success status mismatch', excel_success=success,
                                       matrix_converged=ref['converged'], matrix_status=ref['status']))
                if success != (ref['status'] == 'success'):
                    difference = dict(key=key, issue='matrix status differs from Excel display',
                                      excel_success=success, matrix_converged=ref['converged'], matrix_status=ref['status'])
                    if key == approved_key and success and ref['converged'] and ref['status'] == 'failure':
                        acknowledged.append(difference)
                    else:
                        errors.append(difference)
                if success:
                    for field, value in zip(('beta_hat', 'eta_hat', 'gamma_hat'), values):
                        compared += 1
                        if not same_number(value, ref[field]):
                            errors.append(dict(key=key, issue='estimate value mismatch', field=field,
                                               excel_value=value, matrix_value=ref[field]))
                records.append(dict(combination=cfg['combination'], n=n, id=id_, method=method['label'],
                                    success=success, values=values if success else None))
        sample_rows = list(wb[f'生成样本_n{n}'].iter_rows(values_only=True))
        assert len(sample_rows) == 51 and [row[0] for row in sample_rows[1:]] == list(range(1, 51))
        for row in sample_rows[1:]:
            id_ = int(row[0]); values = list(row[1:n+1]); reference = samples_matrix[(n, id_)]
            assert len(values) == len(reference) == n
            for i, (value, expected) in enumerate(zip(values, reference), 1):
                sample_cells += 1
                if not same_number(value, expected):
                    errors.append(dict(key=[n,id_,i], issue='sample value mismatch', excel_value=value, matrix_value=expected))
            samples.append(dict(n=n, id=id_, values=values))
        for method in cfg['methods']:
            counts.append(dict(n=n, method=method['label'], total=50,
                               success=sum(r['success'] for r in records if r['n']==n and r['method']==method['label'])))
    wb.close()
    report = dict(workbook=str(workbook), workbook_sha256=sha(workbook),
                  crosscheck_matrix=str(matrix_path), matrix_sha256=sha(matrix_path),
                  sample_groups=150, sample_values_compared=sample_cells, method_records=1200,
                  successful_parameter_cells_compared=compared, success_rule='all three Excel cells finite numbers; no parameter trimming',
                  success_matches_matrix=not errors and not acknowledged, comparison_passed=not errors,
                  acknowledged_status_differences=acknowledged, user_confirmed_rule=cfg['success_rule'],
                  boundary_note=cfg['boundary_note'], numeric_comparison='exact float equality, no tolerance',
                  errors=errors, counts=counts)
    (out/'数据一致性核对.json').write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    if errors: raise RuntimeError(f'{len(errors)} Excel/matrix mismatches; stop before statistics or figures')
    payload = dict(source=report, samples=samples, records=records)
    (out/'Excel提取.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    print('EXCEL_MATRIX_VALUE_MATCH', sample_cells, 'sample values;', compared, 'parameter values;',
          'acknowledged status difference', len(acknowledged))
    for row in counts: print(row['n'], row['method'], f"{row['success']}/50")

if __name__=='__main__': main()
