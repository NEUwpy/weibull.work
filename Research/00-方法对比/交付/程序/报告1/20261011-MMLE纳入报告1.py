"""MMLE on Report 1's saved samples: compute once, then replay saved figures.

build: official registry/runner fits; no sampling or replacement of old fits.
render: cache-only plot replay. verify: independent score/root and source checks.
MMLE fixes gamma at the original minimum and deletes one observation. The
conditional curves hold that retained sample fixed; they are not a gamma search.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import sys
import os
import subprocess
import tempfile
import zipfile
import copy
import io
from xml.etree import ElementTree as ET
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
from 共用 import ROOT, RESULT, setup_runtime, load_conditional_profiles, save_png, write_atomic
PROJECT = ROOT.parents[1]
sys.path.insert(0, str(PROJECT / 'python'))
setup_runtime()
import numpy as np

CACHE = HERE / '20261011-MMLE估计与曲线.json'
QA = HERE / '20261011-MMLE核验.json'
DATA = RESULT / '报告1' / '中间数据'
MATRIX = HERE.parent / 'W5,1000,500/完整矩阵/数据/matrix.json'
ORIGINAL = HERE / '输入快照/原案例核验.json'
PAIRED = DATA / '样本.json'
MMLE_CODE = PROJECT / 'python/methods/mmle.py'
FROZEN_MMLE = HERE.parent / 'W5,1000,500/依赖快照/python/methods/mmle.py'
CONDITIONS = [(2., 7), (5., 7), (2., 15), (5., 15)]
BLUE, RETURN, COLOR = '#345D7E', '#B55E3E', '#71896A'


def read(path):
    return json.loads(path.read_text(encoding='utf-8'))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write(path, obj):
    write_atomic(path, (json.dumps(obj, ensure_ascii=False, allow_nan=False,
                                  separators=(',', ':')) + '\n').encode('utf-8'))


def qs(values):
    return np.quantile(values, [0, .25, .5, .75, 1]).tolist()


def fit_sample(s):
    from studies.common.runner import run_method
    r = run_method('mmle', s['observations'])
    return dict(beta=float(s.get('beta', s.get('beta_true'))), n=s['n'],
                sample_id=s['sample_id'], method='mmle',
                beta_hat=r['beta_hat'], eta_hat=r['eta_hat'], gamma_hat=r['gamma_hat'],
                converged=r['converged'], solver=r['extra'])


def source_hashes():
    paths = [ORIGINAL, PAIRED, MATRIX, MMLE_CODE, FROZEN_MMLE,
             PROJECT / 'python/methods/conditional_profiles.py',
             HERE / '输入快照/原案例估计.json', HERE / '图1n30补充.json',
             DATA / '实际估计.json', DATA / '汇总.json']
    return {p.relative_to(PROJECT).as_posix(): sha(p) for p in paths}


def build():
    assert sha(MMLE_CODE) == sha(FROZEN_MMLE), 'MMLE method identity changed'
    sources = source_hashes()
    old = [s for s in read(ORIGINAL)['samples'] if s['beta_true'] == 5]
    pairs = read(PAIRED)
    matrix = read(MATRIX)
    assert read(HERE / '图1n30补充.json')['source_matrix_sha256'] == sha(MATRIX)
    assert len(old) == 100 and len(pairs) == 700
    original_records = [fit_sample(s) for s in old]
    for r in matrix['results']:
        if r['n'] == 30 and r['method_id'] == 'mmle':
            original_records.append(dict(beta=5., n=30, sample_id=r['id'], method='mmle',
                **{k:r[k] for k in ('beta_hat', 'eta_hat', 'gamma_hat', 'converged')}))
    assert len(original_records) == 150
    profiles = load_conditional_profiles()
    paired_records, summaries, curves = [], [], []
    for s in pairs:
        r = fit_sample(s)
        retained = np.sort(s['observations'])[1:]
        truth = profiles.mle_profile(retained, 500., shape_lower=0.)
        r['fixed_truth'] = truth
        paired_records.append(r)
        if (s['beta'], s['n']) in CONDITIONS:
            g = float(min(s['observations']))
            # Insert true location and exact MMLE endpoint as calculated nodes.
            grid = np.unique(np.r_[np.linspace(0., g, 121), 500., g])
            points = []
            for candidate in grid:
                p = profiles.mle_profile(retained, float(candidate), shape_lower=0.)
                assert p is not None
                points.append([float(candidate), p['beta'], p['eta']])
            curves.append(dict(beta=s['beta'], n=s['n'], sample_id=s['sample_id'],
                observations=s['observations'], retained=retained.tolist(),
                fit=r, points=points))
    for beta in sorted({s['beta'] for s in pairs}):
        for n in (7, 15):
            rr = [r for r in paired_records if r['beta'] == beta and r['n'] == n]
            ok = [r for r in rr if r['converged']]
            ss = dict(beta=beta, n=n, method='mmle', total=len(rr), success=len(ok),
                failed=len(rr)-len(ok))
            for k in ('beta_hat', 'eta_hat', 'gamma_hat'):
                ss[k+'_q'] = qs([r[k] for r in ok])
            ss['eta_error_free_q'] = qs([abs(r['eta_hat']-1000) for r in ok])
            ss['eta_error_fixed_q'] = qs([abs(r['fixed_truth']['eta']-1000) for r in ok])
            summaries.append(ss)
    high = [c for c in curves if c['beta'] == 5 and c['n'] == 7 and c['fit']['converged']]
    median_g = float(np.median([c['fit']['gamma_hat'] for c in high]))
    selected_id = min(high, key=lambda c:(abs(c['fit']['gamma_hat']-median_g), c['sample_id']))['sample_id']
    cache = dict(date='2026-10-11', sources=sources, method='Kundu-Raqab single-sample min/delete-one',
        new_samples=0, new_estimates=800, reused_n30_estimates=50,
        original_records=original_records, paired_records=paired_records, summary=summaries,
        curves=curves, selected_id=selected_id,
        selection='n=7 beta=5 gamma nearest median; same id for beta=2; literal expansions',
        diagnostic='Fix original deletion set t2..tn; unrestricted conditional shape root; gamma <= original t1',
        axes=dict(xlim=[-25.,1400.], xticks=[0.,500.,1000.],
                  ylim=[0.,10.], yticks=[0.,2.,4.,6.,8.,10.]))
    assert source_hashes() == sources
    write(CACHE, cache)
    verify()
    print('MMLE saved: 800 new estimates on existing samples; 50 n=30 estimates reused.')


def render_mechanism():
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    from matplotlib.ticker import NullLocator
    c = read(CACHE)
    plt.rcParams.update({'font.family':'sans-serif',
        'font.sans-serif':['Microsoft YaHei','Arial','DejaVu Sans'], 'font.size':7,
        'axes.labelsize':7, 'axes.titlesize':7.2, 'xtick.labelsize':6.5,
        'ytick.labelsize':6.5, 'axes.linewidth':.7, 'axes.spines.top':False,
        'axes.spines.right':False, 'axes.unicode_minus':False, 'legend.frameon':False})
    fig, axes = plt.subplots(3, 2, figsize=(7.2, 6.9))

    def draw(ax, row, strong=False, nodes=False):
        points = np.array(row['points'])
        ax.plot(points[:,0], points[:,1], color=BLUE, alpha=1 if strong else .16,
                lw=1.25 if strong else .65, zorder=3 if strong else 1)
        actual = row['fit']
        if actual['converged']:
            ax.scatter(actual['gamma_hat'], actual['beta_hat'], color=RETURN,
                       marker='x', s=26 if strong else 9,
                       lw=.9 if strong else .5, alpha=1 if strong else .55, zorder=5)
        if strong:
            p = next(p for p in row['points'] if p[0] == 500.)
            ax.scatter(500.,p[1],s=15, facecolor='white',edgecolor=BLUE,lw=.7,zorder=5)
        if nodes:
            ids = np.unique(np.r_[np.arange(0,len(points),12),len(points)-1])
            ax.scatter(points[ids,0],points[ids,1],s=11,facecolor='white',edgecolor=BLUE,lw=.5,zorder=4)
            # Numbered original calculation nodes, repeated identically for both beta examples.
            for label, fraction in enumerate((.2,.55,.85),1):
                idx = int(round(fraction*(len(points)-1)))
                if points[idx,1] <= 9.6:
                    ax.scatter(*points[idx,:2],s=11,facecolor='white',edgecolor=BLUE,lw=.5,zorder=4)
                    ax.annotate(str(label),points[idx,:2],xytext=(0,6),
                                textcoords='offset points',ha='center',fontsize=6,color=BLUE)

    for i,(beta,n) in enumerate(CONDITIONS):
        ax = axes.flat[i]
        rows = [r for r in c['curves'] if (r['beta'],r['n']) == (beta,n)]
        for r in rows:
            draw(ax,r)
        if n == 7:
            chosen = next(r for r in rows if r['sample_id'] == c['selected_id'])
            draw(ax,chosen,strong=True)
        ax.set_title(f'β={int(beta)}, n={n}',loc='left',pad=8)
    for col,beta in enumerate((2.,5.)):
        ax = axes[2,col]
        row = next(r for r in c['curves'] if r['beta'] == beta and r['n'] == 7
                   and r['sample_id'] == c['selected_id'])
        draw(ax,row,strong=True,nodes=True)
        ax.set_title(f'β={int(beta)}, n=7, #{c["selected_id"]}',loc='left',pad=8)
    for i,ax in enumerate(axes.flat):
        ax.set_xlim(c['axes']['xlim']); ax.set_ylim(c['axes']['ylim'])
        ax.set_xticks(c['axes']['xticks']); ax.set_yticks(c['axes']['yticks'])
        ax.xaxis.set_minor_locator(NullLocator()); ax.yaxis.set_minor_locator(NullLocator())
        ax.set_xlabel('γ'); ax.set_ylabel(r'$\hat\beta(\gamma)$')
        ax.axvline(500,color='black',ls=':',lw=.75,zorder=0)
        ax.text(-.12,1.07,chr(97+i),transform=ax.transAxes,weight='bold',fontsize=8)
    fig.legend([Line2D([],[],color=BLUE,lw=1.1),
                Line2D([],[],color=RETURN,marker='x',ls='None',ms=4.5),
                Line2D([],[],color='black',ls=':',lw=.75)],
               ['条件形状','MMLE返回值','真γ'], ncol=3, loc='upper center',
               bbox_to_anchor=(.52,.995),fontsize=7)
    fig.subplots_adjust(left=.10,right=.98,bottom=.07,top=.92,hspace=.58,wspace=.30)
    path = RESULT / '报告1/图8_MMLE六格机制.png'
    save_png(fig,path); plt.close(fig)
    print(f'EXPORTED MMLE six-panel mechanism; paired sample #{c["selected_id"]}.')


def render():
    sys.path.insert(0, str(HERE))
    render_mechanism()
    import 绘制图1小提琴 as violin
    import 绘制补充图 as supplement
    violin.main()
    supplement.continuous_estimates()
    supplement.compensation()
    supplement.supplement(include_lre=False)


def workbook(reuse_export=False):
    """Artifact-tool authors MMLE cells; retain old workbook parts byte for byte.

    Import/export changes old style metadata, so only the authored MMLE sheet
    is attached. Its styles reuse the old header, text and numeric styles.
    No existing sheet, formulas, cached results or style table is rewritten.
    """
    from openpyxl import load_workbook
    node = Path(os.environ.get('CODEX_DEPENDENCIES_NODE',
        'C:/Users/36089/.cache/codex-runtimes/codex-primary-runtime/dependencies/node')) / 'bin/node.exe'
    if not reuse_export:
        subprocess.run([str(node),str(HERE/'20261011-MMLE统计表.mjs')],check=True)
    target = RESULT / '报告1/高形状参数补偿统计.xlsx'
    scratch = Path(tempfile.gettempdir())/'r00-report1-mmle'
    generated = scratch / target.name
    ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    relns='http://schemas.openxmlformats.org/package/2006/relationships'
    docrel='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    ct='http://schemas.openxmlformats.org/package/2006/content-types'
    ET.register_namespace('',ns); ET.register_namespace('r',docrel)
    with zipfile.ZipFile(target) as z:
        original={p:z.read(p) for p in z.namelist()}
    with zipfile.ZipFile(generated) as z:
        exported={p:z.read(p) for p in z.namelist()}
    wb=ET.fromstring(original['xl/workbook.xml'])
    sheets=wb.find(f'{{{ns}}}sheets')
    relations=ET.fromstring(original['xl/_rels/workbook.xml.rels'])
    content_types=ET.fromstring(original['[Content_Types].xml'])
    existed=next((e for e in sheets if e.attrib['name']=='MMLE'),None)
    if existed is None:
        sheet_id=max(int(s.attrib['sheetId']) for s in sheets)+1
        relation_id='rIdMMLE'
        occupied={e.attrib['Id'] for e in relations}
        while relation_id in occupied: relation_id += '_'
        part=f'xl/worksheets/sheet{sheet_id}.xml'
        ET.SubElement(sheets,f'{{{ns}}}sheet',name='MMLE',sheetId=str(sheet_id),
                      attrib={f'{{{docrel}}}id':relation_id})
        ET.SubElement(relations,f'{{{relns}}}Relationship',Id=relation_id,
            Type=docrel+'/worksheet',Target=f'worksheets/sheet{sheet_id}.xml')
        ET.SubElement(content_types,f'{{{ct}}}Override',PartName='/'+part,
            ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
    else:
        relation_id=existed.attrib[f'{{{docrel}}}id']
        relation=next(e for e in relations if e.attrib['Id']==relation_id)
        part='xl/'+relation.attrib['Target'].lstrip('/')
        if part.startswith('xl/xl/'): part=part[3:]
    imported_wb=ET.fromstring(exported['xl/workbook.xml'])
    mmle_sheet=next(s for s in imported_wb.find(f'{{{ns}}}sheets') if s.attrib['name']=='MMLE')
    imported_rels=ET.fromstring(exported['xl/_rels/workbook.xml.rels'])
    source_rel=next(e for e in imported_rels if e.attrib['Id']==mmle_sheet.attrib[f'{{{docrel}}}id'])
    source_part=source_rel.attrib['Target'].lstrip('/')
    if not source_part.startswith('xl/'): source_part='xl/'+source_part
    authored=ET.fromstring(exported[source_part])
    strings=ET.fromstring(exported['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in exported else []
    original_sheet=ET.fromstring(original['xl/worksheets/sheet1.xml'])
    styles={e.attrib['r']:e.attrib.get('s','0') for e in original_sheet.iter(f'{{{ns}}}c')}
    for cell in authored.iter(f'{{{ns}}}c'):
        address=cell.attrib['r']; col=''.join(ch for ch in address if ch.isalpha())
        row=int(''.join(ch for ch in address if ch.isdigit()))
        if cell.attrib.get('t')=='s':
            value=cell.find(f'{{{ns}}}v')
            si=copy.deepcopy(strings[int(value.text)])
            cell.remove(value); si.tag=f'{{{ns}}}is'; cell.append(si);cell.set('t','inlineStr')
        if row in (1,22): style=styles['A1']
        elif cell.attrib.get('t') in ('inlineStr','str'): style=styles['B2']
        elif (row<22 and col in ('C','D')) or (row>22 and col in ('C','D')): style=styles['A2']
        else: style=styles['D2']
        cell.set('s',style)
    updated=dict(original)
    updated[part]=ET.tostring(authored,encoding='utf-8',xml_declaration=True)
    updated['xl/workbook.xml']=ET.tostring(wb,encoding='utf-8',xml_declaration=True)
    updated['xl/_rels/workbook.xml.rels']=ET.tostring(relations,encoding='utf-8',xml_declaration=True)
    updated['[Content_Types].xml']=ET.tostring(content_types,encoding='utf-8',xml_declaration=True)
    allowed={part,'xl/workbook.xml','xl/_rels/workbook.xml.rels','[Content_Types].xml'}
    assert all(updated[k]==v for k,v in original.items() if k not in allowed)
    result=io.BytesIO()
    with zipfile.ZipFile(result,'w',zipfile.ZIP_DEFLATED) as z:
        for name,value in updated.items(): z.writestr(name,value)
    result.seek(0)
    checked=load_workbook(result,read_only=True,data_only=True)
    sh=checked['MMLE']; cached=read(CACHE)
    expected=cached['original_records']+cached['paired_records']
    for row,r in zip(sh.iter_rows(min_row=23,max_row=872,values_only=True),expected):
        assert row[1:4] == (r['beta'],r['n'],r['sample_id'])
        for actual,k in zip(row[5:8],('beta_hat','eta_hat','gamma_hat')):
            assert abs(actual-r[k]) < 1e-9
        assert abs(row[8]-(r['gamma_hat']-500)) < 1e-9
        assert abs(row[9]-(r['eta_hat']-1000)) < 1e-9
    checked.close()
    temporary=target.with_suffix('.tmp.xlsx')
    temporary.write_bytes(result.getvalue());temporary.replace(target)
    previous=read(QA)
    previous['workbook']=dict(path=target.relative_to(ROOT).as_posix(),sha256=sha(target),
        summary_rows=17,detail_rows=850,existing_six_sheet_parts_unchanged=True,
        existing_style_table_unchanged=True,details_checked_against_cache=True,
        authored_by='artifact-tool',added_sheet='MMLE')
    write(QA,previous)
    print('VERIFIED workbook: 850 MMLE records; original six worksheets and styles retained.')


def verify():
    c = read(CACHE)
    assert c['sources'] == source_hashes(), 'Source or method files changed'
    profiles = load_conditional_profiles()
    pairs = {(s['beta'],s['n'],s['sample_id']):s for s in read(PAIRED)}
    originals = {(5.,s['n'],s['sample_id']):s for s in read(ORIGINAL)['samples'] if s['beta_true']==5}
    matrix = read(MATRIX)
    originals.update({(5.,s['n'],s['id']):dict(observations=s['values'])
                      for s in matrix['samples'] if s['n']==30})
    root_errors, eta_errors, scores, node_errors = [], [], [], []
    for group,lookup in ((c['paired_records'],pairs),(c['original_records'],originals)):
        for r in group:
            if not r['converged']:
                continue
            x = sorted(lookup[r['beta'],r['n'],r['sample_id']]['observations'])
            assert r['gamma_hat'] == x[0] and x[1] > r['gamma_hat']
            p = profiles.mle_profile(x[1:],x[0],shape_lower=0.)
            root_errors.append(abs(r['beta_hat']-p['beta']))
            eta_errors.append(abs(r['eta_hat']-p['eta']))
            logs = np.log(np.array(x[1:])-x[0]); b=r['beta_hat']
            w=np.exp(b*(logs-logs.max())); w/=w.sum()
            scores.append(abs(1/b+logs.mean()-w@logs))
    assert max(root_errors) < 5e-6 and max(eta_errors) < 1e-3 and max(scores) < 1e-6
    decreases = 0
    for row in c['curves']:
        pts=np.array(row['points'])
        assert pts[0,0] == 0. and pts[-1,0] == row['observations'][0]
        assert row['retained'] == row['observations'][1:]
        assert 500. in pts[:,0]
        if np.all(np.diff(pts[:,1]) <= 1e-9): decreases += 1
        for idx in (0,len(pts)//2,len(pts)-1):
            p=profiles.mle_profile(row['retained'],float(pts[idx,0]),shape_lower=0.)
            node_errors.append(max(abs(p['beta']-pts[idx,1]),abs(p['eta']-pts[idx,2])))
    assert max(node_errors) < 1e-8
    n30 = {r['id']:r for r in matrix['results'] if r['n']==30 and r['method_id']=='mmle'}
    for r in c['original_records']:
        if r['n']==30:
            assert all(r[k] == n30[r['sample_id']][k] for k in ('beta_hat','eta_hat','gamma_hat','converged'))
    out=dict(date='2026-10-11', cache_sha256=sha(CACHE), inputs_unchanged=True,
        new_samples=0, successful=sum(r['converged'] for r in c['paired_records']+c['original_records']),
        total=850, original_case_records=150, paired_records=700, curves=len(c['curves']),
        conditional_shape_decreasing_curves=decreases,
        maximum_shape_root_error=max(root_errors), maximum_scale_root_error=max(eta_errors),
        maximum_shape_score_residual=max(scores), maximum_checked_node_error=max(node_errors),
        n30_exactly_matches_preserved_batch=True, source_hashes=c['sources'],
        single_examples_are_literal_expansions=True, selected_id=c['selected_id'],
        formats=['PNG'], common_linear_axes=c['axes'])
    if QA.exists():
        previous=read(QA)
        for key in ('workbook','delivery'):
            if key in previous: out[key]=previous[key]
        if 'workbook' in out:
            assert sha(ROOT/out['workbook']['path']) == out['workbook']['sha256']
    write(QA,out)
    print(f'VERIFIED {out["successful"]}/{out["total"]} successful MMLE fits, {len(c["curves"])} conditional curves.')


if __name__ == '__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('action',nargs='?',default='render',choices=['build','render','verify','workbook'])
    parser.add_argument('--use-exported',action='store_true',help='Reuse a just-verified temporary workbook export')
    args=parser.parse_args()
    if args.action=='workbook': workbook(reuse_export=args.use_exported)
    else: globals()[args.action]()
