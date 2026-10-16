"""Export the research reports with the supplied template and native MathType."""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image
from lxml import etree
from zipfile import ZipFile, ZIP_DEFLATED

HERE = Path(__file__).resolve().parent
BASE = HERE.parents[2]
SOURCE = BASE / '报告1：小样本情况下高形状参数的估计偏移.md'
OUTPUT = SOURCE.with_suffix('.docx')
REFERENCE = HERE / '20261012-报告1转Word-原模板.docx'
QA = BASE.parents[1] / 'tmp' / 'report1-word-20261012'
REPORT_ID = 1
REPORT_TITLE = '小样本情况下高形状参数的估计偏移'
EXPECTED_IMAGES = 15
EXPECTED_FORMULAS = 34
WRAPPED_FORMULAS = {}
MATHJAX_LIBRARY = QA / 'mathjax/node_modules/mathjax-full'
POWERSHELL = Path(r'C:\Users\36089\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\powershell\pwsh.exe')
NODE = Path(r'C:\Users\36089\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\bin\node.exe')
sys.path.insert(0, str(QA/'deps'))


def plain(text):
    return re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)


def prepare():
    QA.mkdir(parents=True, exist_ok=True)
    try:
        from latex2mathml.converter import convert as latex_to_mathml
    except ImportError:
        dependency_dir = BASE.parents[1]/'tmp/report1-word-20261012/deps'
        subprocess.run([sys.executable, '-m', 'pip', 'install', '--target', str(dependency_dir),
                        'latex2mathml==3.81.1', '--quiet', '--disable-pip-version-check'], check=True)
        sys.path.insert(0, str(dependency_dir))
        from latex2mathml.converter import convert as latex_to_mathml
    if not REFERENCE.exists():
        backup = QA / 'template-original.docx'
        REFERENCE.write_bytes(backup.read_bytes() if backup.exists() else OUTPUT.read_bytes())
    doc = Document(REFERENCE)
    source = SOURCE.read_text(encoding='utf-8')
    # Reuse the actual main-body section, heading formatting, styles, and footer.
    paragraphs = list(doc.paragraphs)
    heading1 = next(p for p in paragraphs if p.style.style_id == '1')
    heading2 = next(p for p in paragraphs if p.style.style_id == '2')
    h1p, h1r = deepcopy(heading1._p.pPr), deepcopy(heading1.runs[0]._r.rPr)
    h2p, h2r = deepcopy(heading2._p.pPr), deepcopy(heading2.runs[0]._r.rPr)
    bp, br = deepcopy(paragraphs[28]._p.pPr), deepcopy(paragraphs[28].runs[0]._r.rPr)
    cp, cr = deepcopy(paragraphs[110]._p.pPr), deepcopy(paragraphs[110].runs[0]._r.rPr)
    section = deepcopy(doc.sections[-1]._sectPr)
    body = doc._element.body
    for child in list(body):
        body.remove(child)
    body.append(section)
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith(('/image', '/oleObject', '/hyperlink')):
            doc.part.drop_rel(rid)
    doc.core_properties.title = REPORT_TITLE
    doc.core_properties.subject = f'报告{REPORT_ID}'
    width_pt = (doc.sections[0].page_width - doc.sections[0].left_margin
                - doc.sections[0].right_margin) / 12700
    jobs = []
    source_paragraphs = []
    tables = []

    def add_inline(paragraph, text, run_properties):
        for token in re.split(r'(\$[^$]+\$|\*\*.*?\*\*)', plain(text)):
            if not token:
                continue
            if token.startswith('$') and token.endswith('$'):
                marker = f'__R{REPORT_ID}EQ{len(jobs)+1:04d}__'
                paragraph.add_run(marker)
                jobs.append({'marker': marker, 'tex': token[1:-1], 'display': False,
                             'number': None, 'max_width_pt': width_pt})
            else:
                run = paragraph.add_run(token[2:-2] if token.startswith('**') else token)
                run._r.insert(0, deepcopy(run_properties))
                run.bold = token.startswith('**')

    def display(tex):
        match = re.search(r'\\tag\{(\d+)\}', tex)
        number = int(match[1])
        clean = re.sub(r'\\tag\{\d+\}', '', tex).strip()
        if number in WRAPPED_FORMULAS:
            clean = WRAPPED_FORMULAS[number]
        # Wrap only the presentation of the two long implication formulas.
        if REPORT_ID == 1 and number in (26, 28):
            left, right = clean.split(r'\quad\Longrightarrow\quad')
            left = left.replace('=', '&=', 1)
            clean = r'\begin{aligned}' + left.strip() + r'\\ &\Longrightarrow\quad ' + right.strip() + r'\end{aligned}'
        if REPORT_ID == 1 and number == 31:
            clean = clean.replace(r'\right),\qquad', r'\right),\\')
            clean = clean.replace(r'z_i=\frac', r'z_i&=\frac')
        clean = clean.replace(r'\begin{aligned}', r'\begin{array}{rl}').replace(r'\end{aligned}', r'\end{array}')
        paragraph = doc.add_paragraph(style='MTDisplayEquation')
        pf = paragraph.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.left_indent = Pt(0)
        pf.right_indent = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.space_before, pf.space_after = Pt(5), Pt(7)
        pf.keep_together = True
        pf.keep_with_next = REPORT_ID == 1
        props = paragraph._p.get_or_add_pPr()
        for name in ('ind', 'tabs'):
            old = props.find(qn('w:' + name))
            if old is not None:
                props.remove(old)
        tabs = OxmlElement('w:tabs')
        for value, pos in [('clear', 7140), ('center', int(width_pt*10)), ('right', int(width_pt*20))]:
            tab = OxmlElement('w:tab')
            tab.set(qn('w:val'), value)
            tab.set(qn('w:pos'), str(pos))
            tabs.append(tab)
        props.append(tabs)
        marker = f'__R{REPORT_ID}EQ{len(jobs)+1:04d}__'
        paragraph.add_run('\t' + marker + f'\t({number})')
        jobs.append({'marker': marker, 'tex': clean, 'source_tex': tex.strip(),
                     'display': True, 'number': number, 'max_width_pt': width_pt-(100 if REPORT_ID == 1 and number == 27 else 48)})

    def add_table(text):
        rows = [[v.strip() for v in line.strip().strip('|').split('|')]
                for line in text.splitlines() if line.strip()]
        assert all(re.fullmatch(r':?-+:?', v.replace(' ', '')) for v in rows[1])
        rows.pop(1)
        count = len(rows[0])
        assert all(len(row) == count for row in rows)
        # Table typography follows the reference's actual data table (10.5 pt).
        if count == 2:
            widths = [65, width_pt-65]
        elif rows[0][-1].startswith('观测值') and count == 5:
            widths = [70, 45, 60, 60, width_pt-235]
        elif rows[0][-1].startswith('观测值') and count == 7:
            widths = [66, 72, 20, 36, 47, 47, width_pt-288]
        else:
            widths = [width_pt/count]*count
        table = doc.add_table(rows=len(rows), cols=count)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        margins = OxmlElement('w:tblCellMar')
        for side, value in [('top', 65), ('bottom', 65), ('left', 80), ('right', 80)]:
            node = OxmlElement('w:'+side)
            node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')
            margins.append(node)
        table._tbl.tblPr.append(margins)
        for column, value in zip(table.columns, widths):
            column.width = Pt(value)
        for row_index, values in enumerate(rows):
            trpr = table.rows[row_index]._tr.get_or_add_trPr()
            trpr.append(OxmlElement('w:cantSplit'))
            if row_index == 0:
                trpr.append(OxmlElement('w:tblHeader'))
            for column_index, (cell, value, width) in enumerate(zip(table.rows[row_index].cells, values, widths)):
                cell.width = Pt(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph = cell.paragraphs[0]
                # Preserve table typography without the body indent or line grid.
                snap = OxmlElement('w:snapToGrid')
                snap.set(qn('w:val'), '0')
                paragraph._p.get_or_add_pPr().append(snap)
                pf = paragraph.paragraph_format
                pf.first_line_indent = Pt(0)
                pf.space_before, pf.space_after = Pt(2), Pt(2)
                pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                pf.keep_with_next = row_index < len(rows)-1
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if count == 2 or column_index == count-1 and rows[0][-1].startswith('观测值') else WD_ALIGN_PARAGRAPH.CENTER
                shown = value
                if rows[0][column_index] == '参数组合' and len(value) > 13 and value.startswith('W('):
                    first, last = value.rsplit(',', 1)
                    shown = first+',\n'+last
                add_inline(paragraph, shown, br)
                for run in paragraph.runs:
                    run.font.size = Pt(10.5)
                    if row_index == 0:
                        run.bold = True
        tables.append(rows)

    blocks = re.split(r'(\$\$.*?\$\$)', source, flags=re.S)
    figures = []
    for block in blocks:
        if block.startswith('$$'):
            display(block[2:-2])
            continue
        for text in re.split(r'\n\s*\n', block.strip()):
            text = text.strip()
            if not text:
                continue
            # Report title/date are metadata; the requested document starts at chapter 1.
            if text.startswith('# ') or text.startswith('> '):
                continue
            if text.startswith('|'):
                add_table(text)
                continue
            if text.startswith('#'):
                level = len(text)-len(text.lstrip('#'))
                paragraph = doc.add_paragraph()
                paragraph._p.append(deepcopy(h1p if level == 2 else h2p))
                run = paragraph.add_run(plain(text.lstrip('#').strip()))
                run._r.insert(0, deepcopy(h1r if level == 2 else h2r))
                if level >= 4:
                    run.font.size = Pt(12)
                    run.bold = True
                    paragraph.style = 'Heading 3'
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.keep_with_next = True
                if level == 2:
                    paragraph.paragraph_format.page_break_before = True
                continue
            image = re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)', text)
            if image:
                path = SOURCE.parent / image[2]
                iw, ih = Image.open(path).size
                height_limit = 560 if image[1].startswith('图2　') else 410
                width = min(width_pt, height_limit*iw/ih)
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = paragraph.paragraph_format
                pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                pf.first_line_indent = Pt(0)
                pf.space_before, pf.space_after = Pt(6), Pt(3)
                pf.keep_with_next, pf.keep_together = True, True
                paragraph.add_run().add_picture(str(path), width=Pt(width))
                figures.append({'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                                'title': image[1], 'width_pt': width, 'height_pt': width*ih/iw})
                continue
            caption = text.startswith('**图')
            paragraph = doc.add_paragraph()
            paragraph._p.append(deepcopy(cp if caption else bp))
            if not caption:
                paragraph.paragraph_format.widow_control = True
                if text.endswith(('：', ':')):
                    paragraph.paragraph_format.keep_with_next = True
                if REPORT_ID == 2 and text.startswith('**') and text.endswith('**'):
                    paragraph.paragraph_format.keep_with_next = True
            if caption:
                pf = paragraph.paragraph_format
                pf.keep_with_next = False
                pf.keep_together = True
                pf.first_line_indent = Pt(0)
                pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                pf.space_after = Pt(8)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, text.replace('\n', ''), cr if caption else br)
            source_paragraphs.append(plain(text.replace('\n', '')))
    assert len(figures) == EXPECTED_IMAGES
    assert [j['number'] for j in jobs if j['display']] == list(range(1, EXPECTED_FORMULAS+1))
    for job in jobs:
        job['mathml'] = latex_to_mathml(job['tex'], display='block' if job['display'] else 'inline')
        etree.fromstring(job['mathml'].encode('utf-8'))
    # Save a draft; the native Word/MathType stage replaces every marker.
    stem = f'report{REPORT_ID}'
    draft = QA / f'{stem}-draft.docx'
    doc.save(draft)
    manifest = {'source': str(SOURCE), 'source_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
                'reference': str(REFERENCE), 'reference_sha256': hashlib.sha256(REFERENCE.read_bytes()).hexdigest(),
                'draft': str(draft), 'output': str(OUTPUT), 'qa': str(QA),
                'stem': stem, 'mathjax_library': str(MATHJAX_LIBRARY),
                'width_pt': width_pt, 'jobs': jobs, 'figures': figures,
                'source_paragraphs': source_paragraphs, 'tables': tables}
    path = QA / 'manifest.json'
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    (QA / 'artifact.md').write_text(
        f'# 报告{REPORT_ID} Word 模板规范\n\n'
        f'参考：{REFERENCE}\n\nA4；上下2.54 cm，左右3.175 cm。'
        '沿用主正文节、页脚页码和全部原有样式；移除封面、目录及原正文。'
        '一级标题沿用模板正文的15 pt，二级标题14 pt；正文沿用模板实际正文的12 pt仿宋、28 pt行距。'
        '图居中、等比例缩放，图题沿用模板实际图题的12 pt仿宋。'
        '编号公式使用MTDisplayEquation样式、MathType OLE对象和右侧原编号；'
        f'公式及图片段落改用自动行距以容纳对象。Markdown正文、公式和{EXPECTED_IMAGES}张源图保留。'
        '数据表格采用模板的Table Grid样式、10.5 pt字体、自动行高及重复表头。\n',
        encoding='utf-8')
    print(f'Prepared {len(figures)} figures, {EXPECTED_FORMULAS} numbered equations, {len(jobs)-EXPECTED_FORMULAS} inline equations, {len(tables)} tables.')
    return path


def compose(manifest):
    config = json.loads(manifest.read_text(encoding='utf-8'))
    metrics = json.loads((QA/'preview-metrics.json').read_text(encoding='utf-8'))
    with ZipFile(config['draft']) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(parts['word/document.xml'])
    rels = etree.fromstring(parts['word/_rels/document.xml.rels'])
    types = etree.fromstring(parts['[Content_Types].xml'])
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
          'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
          'v': 'urn:schemas-microsoft-com:vml', 'o': 'urn:schemas-microsoft-com:office:office'}
    relns = 'http://schemas.openxmlformats.org/package/2006/relationships'
    ctns = 'http://schemas.openxmlformats.org/package/2006/content-types'
    for index, job in enumerate(config['jobs'], 1):
        path = QA / 'objects' / f'{job["marker"]}.docx'
        with ZipFile(path) as archive:
            eqdoc = etree.fromstring(archive.read('word/document.xml'))
            eqrels = {r.get('Id'): r for r in etree.fromstring(archive.read('word/_rels/document.xml.rels'))}
            obj = deepcopy(eqdoc.xpath('//w:object', namespaces=ns)[0])
            dimension = metrics[job['marker']]
            obj.set(qn('w:dxaOrig'), str(round(dimension['width_pt']*20)))
            obj.set(qn('w:dyaOrig'), str(round(dimension['height_pt']*20)))
            shapeid = f'_x0000_i{2000+index}'
            for shape in obj.xpath('.//v:shape', namespaces=ns):
                shape.set('id', shapeid)
                style = shape.get('style','')
                style = re.sub(r'width:[^;]+', f'width:{dimension["width_pt"]:.3f}pt', style)
                style = re.sub(r'height:[^;]+', f'height:{dimension["height_pt"]:.3f}pt', style)
                shape.set('style',style)
            for ole in obj.xpath('.//o:OLEObject', namespaces=ns):
                ole.set('ShapeID', shapeid)
                ole.set('ObjectID', f'_{2000000000+index}')
            for node in obj.iter():
                oldid = node.get(qn('r:id'))
                if not oldid:
                    continue
                rel = eqrels[oldid]
                target = rel.get('Target')
                suffix = Path(target).suffix
                category = 'embeddings' if rel.get('Type').endswith('/oleObject') else 'media'
                data = archive.read('word/'+target)
                if category == 'media' and suffix.lower() == '.wmf':
                    # Replace only the presentation cache; preserve native OLE.
                    data, suffix = (QA/'objects'/(job['marker']+'.png')).read_bytes(), '.png'
                newtarget = f'{category}/{config["stem"]}-math-{index:04d}{suffix}'
                parts['word/'+newtarget] = data
                newid = f'rIdReport{REPORT_ID}Math{index}_{category}'
                node.set(qn('r:id'), newid)
                item = etree.SubElement(rels, '{'+relns+'}Relationship')
                item.attrib.update({'Id': newid, 'Type': rel.get('Type'), 'Target': newtarget})
            eqtypes = etree.fromstring(archive.read('[Content_Types].xml'))
            for item in eqtypes:
                ext = item.get('Extension')
                if ext and not types.xpath('ct:Default[@Extension="'+ext+'"]', namespaces={'ct': ctns}):
                    types.append(deepcopy(item))
        markers = root.xpath('//w:t[text()="'+job['marker']+'"]', namespaces=ns)
        assert len(markers) == 1, job['marker']
        marker = markers[0]
        marker.getparent().replace(marker, obj)
    assert len(root.xpath('//o:OLEObject[@ProgID="Equation.DSMT4"]', namespaces=ns)) == len(config['jobs'])
    assert not root.xpath('//w:t[contains(text(),"__R'+str(REPORT_ID)+'EQ")]', namespaces=ns)
    for name, value in [('word/document.xml', root), ('word/_rels/document.xml.rels', rels), ('[Content_Types].xml', types)]:
        parts[name] = etree.tostring(value, encoding='UTF-8', xml_declaration=True, standalone=True)
    with ZipFile(QA/f'{config["stem"]}-complete.docx', 'w', ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            archive.writestr(name, value)
    print(f'Composed {len(config["jobs"])} editable MathType objects and {len(config["figures"])} figures.')


def run(prepare_only=False):
    manifest = prepare()
    if prepare_only:
        return
    if not (MATHJAX_LIBRARY/'js/mathjax.js').is_file():
        pnpm = NODE.parents[2]/'bin/fallback/pnpm.cmd'
        MATHJAX_LIBRARY.parents[1].mkdir(parents=True, exist_ok=True)
        subprocess.run([str(pnpm), '--dir', str(MATHJAX_LIBRARY.parents[1]),
                        'add', 'mathjax-full@3.2.2'], check=True)
    subprocess.run([str(POWERSHELL), '-NoProfile', '-ExecutionPolicy', 'Bypass',
                    '-File', str(HERE/'20261012-报告1转Word-MathType.ps1'),
                    '-Manifest', str(manifest)], check=True)
    subprocess.run([str(NODE), str(HERE/'20261012-报告1转Word-公式预览.cjs'), str(manifest)], check=True)
    compose(manifest)
    subprocess.run([str(POWERSHELL), '-NoProfile', '-ExecutionPolicy', 'Bypass',
                    '-File', str(HERE/'20261012-报告1转Word-MathType.ps1'),
                    '-Manifest', str(manifest), '-RenderOnly'], check=True)
    OUTPUT.write_bytes((QA/f'report{REPORT_ID}-complete.docx').read_bytes())


def verify():
    """Check source coverage, template fidelity and untouched native equation data."""
    config = json.loads((QA/'manifest.json').read_text(encoding='utf-8'))
    config.setdefault('stem', f'report{REPORT_ID}')
    config.setdefault('tables', [])
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
          'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
          'o': 'urn:schemas-microsoft-com:office:office'}
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == config['source_sha256']
    assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == config['reference_sha256']
    with ZipFile(OUTPUT) as final, ZipFile(REFERENCE) as reference:
        root = etree.fromstring(final.read('word/document.xml'))
        rels = {r.get('Id'): r for r in etree.fromstring(final.read('word/_rels/document.xml.rels'))}
        objects = root.xpath('//o:OLEObject', namespaces=ns)
        assert len(objects) == len(config['jobs'])
        assert all(o.get('ProgID') == 'Equation.DSMT4' for o in objects)
        replacements = {}
        for index, job in enumerate(config['jobs'], 1):
            filename = f'word/embeddings/{config["stem"]}-math-{index:04d}.bin'
            data = final.read(filename)
            assert data[:8] == bytes.fromhex('d0cf11e0a1b11ae1')
            with ZipFile(QA/'objects'/f'{job["marker"]}.docx') as native:
                bins = [n for n in native.namelist() if n.startswith('word/embeddings/')]
                assert len(bins) == 1 and data == native.read(bins[0]), job['marker']
            replacements[filename.removeprefix('word/')] = job
        # Reconstruct inline source formulas for exact paragraph coverage.
        copy = deepcopy(root)
        for obj in copy.xpath('//w:object', namespaces=ns):
            ole = obj.xpath('.//o:OLEObject', namespaces=ns)[0]
            job = replacements[rels[ole.get(qn('r:id'))].get('Target')]
            text = OxmlElement('w:t')
            text.text = '' if job['display'] else '$'+job['tex']+'$'
            obj.getparent().replace(obj, text)
        paragraphs = [''.join(p.xpath('.//w:t/text()', namespaces=ns))
                      for p in copy.xpath('/w:document/w:body/w:p', namespaces=ns)]
        for paragraph in config['source_paragraphs']:
            assert paragraph.replace('**','') in paragraphs, paragraph[:70]
        actual_tables = [[[''.join(c.xpath('.//w:t/text()', namespaces=ns))
                           for c in row.xpath('./w:tc', namespaces=ns)]
                          for row in table.xpath('./w:tr', namespaces=ns)]
                         for table in copy.xpath('/w:document/w:body/w:tbl', namespaces=ns)]
        expected_tables = [[[plain(c).replace('**','') for c in row] for row in table]
                           for table in config['tables']]
        assert actual_tables == expected_tables
        figure_hashes = {hashlib.sha256(final.read(n)).hexdigest()
                         for n in final.namelist() if n.startswith('word/media/')}
        assert all(f['sha256'] in figure_hashes for f in config['figures'])
        assert etree.tostring(etree.fromstring(final.read('word/styles.xml')), method='c14n') == etree.tostring(etree.fromstring(reference.read('word/styles.xml')), method='c14n')
        refroot = etree.fromstring(reference.read('word/document.xml'))
        for name in ('pgSz', 'pgMar', 'docGrid'):
            assert root.xpath('//w:sectPr/w:'+name, namespaces=ns)[-1].attrib == refroot.xpath('//w:sectPr/w:'+name, namespaces=ns)[-1].attrib
        assert not root.xpath('//*[local-name()="oMath"]')
        assert not any('__R' in p or '$' in p for p in [''.join(n.xpath('.//w:t/text()', namespaces=ns)) for n in root.xpath('//w:p', namespaces=ns)])
        assert not root.xpath('//w:instrText[contains(text(),"TOC")]', namespaces=ns)
    result = {'source_sha256': config['source_sha256'],
              'output_sha256': hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
              'figures': len(config['figures']), 'tables': len(config['tables']),
              'numbered_equations': EXPECTED_FORMULAS, 'native_mathtype_objects': len(config['jobs']),
              'source_coverage': True, 'native_data_unchanged': True, 'template_fidelity': True}
    (QA/'final-verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--prepare-only', action='store_true')
    parser.add_argument('--verify', action='store_true')
    args = parser.parse_args()
    verify() if args.verify else run(args.prepare_only)
