import json, sys, zipfile
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from PIL import Image
ROOT=Path('D:/weibull')
WORK=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT/'python'))
from studies.common.sample import generate_sample
from studies.common.runner import run_method
d=json.loads((WORK/'payload.json').read_text(encoding='utf-8'))['cases'][0]
folder=WORK/'outputs'/d['label']
wb=load_workbook(folder/'2,1000,1000.xlsx',data_only=True)
cells=0
for n in (7,15,30):
 for i,s in enumerate(d['samples'][str(n)]):
  assert np.array_equal(s,generate_sample(2.0,1000.0,1000.0,n,i,seed=20260921))
  got=[wb[f'生成样本_n{n}'].cell(i+2,j+2).value for j in range(n)]
  np.testing.assert_allclose(got,s,rtol=1e-14); cells+=n
 for off,start in [('0.10',5),('0.15',63),('0.20',121)]:
  for i,row in enumerate(d['results'][str(n)][off]):
   for j,name in enumerate(['MDM','LS','LRE','WMLM','MLM']):
    r=row[name]
    expected=[r['shape_hat'],r['scale_hat'],r['location_hat']] if r['converged'] else ['—']*3
    actual=[wb[f'估计结果_n{n}'].cell(start+i,2+j*3+k).value for k in range(3)]
    if r['converged']: np.testing.assert_allclose(actual,expected,rtol=1e-14,atol=1e-12)
    else: assert actual==expected
    cells+=3
checks=0
maxdiff=0
for n in (7,15,30):
 for sid in (1,25,50):
  sample=[wb[f'生成样本_n{n}'].cell(sid+1,j+2).value for j in range(n)]
  for mid,name in [('mdm','MDM'),('lse','LS'),('lre','LRE'),('wmle','WMLM'),('mle','MLM')]:
   for off in ([.1,.15,.2] if mid=='mdm' else [.1]):
    # Read the stored full precision sample from payload, and verify Excel precision above.
    out=run_method(mid,d['samples'][str(n)][sid-1],**({'offset':off,'gamma_steps':240} if mid=='mdm' else {}))
    exp=d['results'][str(n)][f'{off:.2f}'][sid-1][name]
    assert out['converged']==exp['converged'],(n,sid,name,'status')
    if out['converged']:
     diff=max(abs(out[k]-exp[v]) for k,v in [('beta_hat','shape_hat'),('eta_hat','scale_hat'),('gamma_hat','location_hat')])
     maxdiff=max(maxdiff,diff)
     assert diff<1e-6,(n,sid,name,diff)
    checks+=1
print('sample_regeneration=150 groups identical; workbook_cells=',cells,'recalculations=',checks,'max_difference=',maxdiff)
print('diagnostics=',json.dumps(d['diagnostics']))
images=sorted(folder.glob('*.png'))
assert len(images)==9
print('images=',len(images),'dimensions=',{Image.open(p).size for p in images})
thumbs=[]
for p in images:
 im=Image.open(p).convert('RGB'); im.thumbnail((600,440)); thumbs.append(im)
contact=Image.new('RGB',(1800,1320),'white')
for i,im in enumerate(thumbs):contact.paste(im,((i%3)*600,(i//3)*440))
contact.save(WORK/'qa'/'plots-contact.png')
