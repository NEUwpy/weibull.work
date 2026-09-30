from pathlib import Path
from copy import deepcopy
import zipfile, re, xml.etree.ElementTree as ET
from openpyxl import load_workbook

root=Path('D:/weibull/docs/临时任务/临时任务-W2-1000-3000-MDM偏移量估计-20260825/260921-W2参数估计案例')
p=root/'W(2,1000,1000)'/'2,1000,1000.xlsx'
before=load_workbook(p)
values={s.title:list(s.values) for s in before}
ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def q(s):return '{'+ns+'}'+s
# The artifact range copy omitted cross-sheet style metadata. Transfer only
# existing OOXML style references and layout, preserving every cell value.
with zipfile.ZipFile(p) as z: parts={n:z.read(n) for n in z.namelist()}
for src,dst,samples in [(2,5,True),(4,6,False)]:
 source=ET.fromstring(parts[f'xl/worksheets/sheet{src}.xml'])
 target=ET.fromstring(parts[f'xl/worksheets/sheet{dst}.xml'])
 sc={c.attrib['r']:c for c in source.findall('.//'+q('c'))}
 sr={r.attrib['r']:r for r in source.findall('.//'+q('row'))}
 for row in target.findall('.//'+q('row')):
  ref=sr.get(row.attrib['r'])
  if ref is not None:
   for k in ('ht','customHeight','s','customFormat','hidden'):
    row.attrib.pop(k,None)
    if k in ref.attrib:row.set(k,ref.attrib[k])
  for c in row.findall(q('c')):
   address=c.attrib['r'];col=re.match('[A-Z]+',address)[0]
   refaddr=('P'+re.search('[0-9]+',address)[0]) if samples and (len(col)>1 or col>'P') else address
   ref=sc.get(refaddr)
   c.attrib.pop('s',None)
   if ref is not None and 's' in ref.attrib:c.set('s',ref.attrib['s'])
 for tag in ('sheetViews','sheetFormatPr','mergeCells','pageMargins','pageSetup','printOptions'):
  old=target.find(q(tag));new=source.find(q(tag))
  idx=list(target).index(old) if old is not None else len(target)
  if old is not None:target.remove(old)
  if new is not None:target.insert(idx,deepcopy(new))
 # Keep XML children in worksheet schema order.
 order=['sheetPr','dimension','sheetViews','sheetFormatPr','cols','sheetData','sheetCalcPr','sheetProtection','protectedRanges','scenarios','autoFilter','sortState','dataConsolidate','customSheetViews','mergeCells','phoneticPr','conditionalFormatting','dataValidations','hyperlinks','printOptions','pageMargins','pageSetup','headerFooter','rowBreaks','colBreaks','customProperties','cellWatches','ignoredErrors','smartTags','drawing','legacyDrawing','legacyDrawingHF','picture','oleObjects','controls','webPublishItems','tableParts','extLst']
 target[:]=sorted(target,key=lambda e:order.index(e.tag.split('}')[-1]) if e.tag.split('}')[-1] in order else 99)
 parts[f'xl/worksheets/sheet{dst}.xml']=ET.tostring(target,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile(p,'w',zipfile.ZIP_DEFLATED) as z:
 for n,data in parts.items():z.writestr(n,data)
after=load_workbook(p)
assert values=={s.title:list(s.values) for s in after},'Values changed'
for src,dst,width in [('生成样本_n15','生成样本_n30',31),('估计结果_n15','估计结果_n30',16)]:
 a,b=after[src],after[dst]
 for row in range(1,a.max_row+1):
  assert a.row_dimensions[row].height==b.row_dimensions[row].height
  for col in range(1,width+1):
   assert a.cell(row,min(col,16))._style==b.cell(row,col)._style,(dst,row,col)
 assert a.sheet_view.showGridLines==b.sheet_view.showGridLines
 if width==16:assert str(a.merged_cells)==str(b.merged_cells)
 print(dst,'styles and row heights match')
with zipfile.ZipFile(root/'W(2,1000,1000).zip','w',zipfile.ZIP_DEFLATED) as z:
 for f in sorted(p.parent.iterdir()):
  if f.suffix in ('.xlsx','.png'):z.write(f,f.name)
print('All six sheets values unchanged; ZIP updated.')
