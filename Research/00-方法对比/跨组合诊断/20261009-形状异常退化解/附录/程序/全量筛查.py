if __name__ == '__main__' and len(__import__('sys').argv) != 2:
 raise SystemExit('需提供独立输出目录；不会修改源实验。')
import sys
from pathlib import Path
import json,csv,math,collections
import openpyxl
O=Path(sys.argv[1]); O.mkdir(parents=True, exist_ok=True);R=Path(r'D:\weibull\Research');R00=R/'00-方法对比'
def write(p,rows):
 with p.open('w',encoding='utf8',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]),lineterminator='\n');w.writeheader();w.writerows(rows)
def finite(x):return isinstance(x,(int,float)) and not isinstance(x,bool) and math.isfinite(x)
def method(r):return f"MDM δ={r['delta']:.2f}" if r['method_id']=='mdm' else r['method_id'].upper()
allrows=[];samples={};audit=[];meta=[]
sources=json.loads((R00/'跨组合汇总/程序/20261007-新增五组合/13组合取数.json').read_text(encoding='utf8'))['sources']
for s in sources:
 p=Path(s['program']);batch=p.parent if p.name=='程序' else p.parents[1]
 xlsx=next((batch/'结果').glob('W*.xlsx'));mat=p/'中间数据/matrix.json' if p.name=='程序' else p.parent/'完整矩阵/数据/matrix.json'
 m=json.loads(mat.read_text(encoding='utf8'));truth=m['config']['truth'];idx={(r['n'],r['id'],r['method_id'],r['delta']):r for r in m['results']}
 cfg=json.loads((p/'config.json').read_text(encoding='utf8'));methods=cfg['methods']
 si={(int(r['n']),int(r['id'])):r['values'] for r in m['samples']}
 wb=openpyxl.load_workbook(xlsx,read_only=True,data_only=True);errors=[];nc=0;samplecells=0;boundary=[]
 for n in [7,15,30]:
  xrows=list(wb[f'生成样本_n{n}'].iter_rows(values_only=True));rr=list(wb[f'估计结果_n{n}'].iter_rows(values_only=True))
  assert len(xrows)==51 and len(rr)==52 and list(rr[1][1:])==['α','β','γ']*8
  for sid in range(1,51):
   x=si[n,sid];assert list(xrows[sid][1:n+1])==x;samplecells+=n
   samples[f"R00|{s['combination']}|{n}|{sid}"]=x
   for spec in methods:
    r=idx[n,sid,spec['method_id'],spec['delta']];pos=rr[0].index(spec['excel_label']);vals=rr[sid+1][pos:pos+3];ok=all(finite(v) for v in vals)
    expected=bool(r['converged'] and all(finite(r[k]) for k in ['beta_hat','eta_hat','gamma_hat']))
    assert expected==ok,(s['combination'],n,sid,spec)
    if ok:
     assert all(float(a)==float(r[k]) for a,k in zip(vals,['beta_hat','eta_hat','gamma_hat']));nc+=3
    if ok and r['status']!='success':boundary.append([n,sid,method(r),r['status']])
    allrows.append(dict(project='R00',combination=s['combination'],n=n,group=sid,method=method(r),lre_version='Park Proposed+Plot' if s['protocol']!='historical-eight' else 'Bernard',beta_truth=truth[0],eta_truth=truth[1],gamma_truth=truth[2],beta_hat=r['beta_hat'],eta_hat=r['eta_hat'],gamma_hat=r['gamma_hat'],sample_min=x[0],converged=bool(r['converged']),status=r['status'],excel_success=ok,workbook=str(xlsx),sheet=f'估计结果_n{n}',range=f'{openpyxl.utils.get_column_letter(pos+1)}{sid+2}:{openpyxl.utils.get_column_letter(pos+3)}{sid+2}'))
 wb.close();audit.append(dict(project='R00',combination=s['combination'],parameter_cells=nc,sample_cells=samplecells,records=1200,unexpected=errors,status_differences=boundary))
 meta.append(dict(s,xlsx=str(xlsx),matrix=str(mat)));print('R00 PASS',s['combination'],flush=True)
assert len(allrows)==15600 and len(samples)==1950
for r in allrows:
 valid=all(finite(r[k]) for k in ['beta_hat','eta_hat','gamma_hat']) and r['beta_hat']>0 and r['eta_hat']>0
 r['high_absolute']=bool(valid and r['beta_hat']>=10)
 r['high_relative']=bool(valid and r['beta_hat']>=5*r['beta_truth'])
 r['location_zero']=bool(valid and abs(r['gamma_hat'])<=1e-6*r['sample_min'])
 r['scale_large']=bool(valid and r['eta_hat']>=5*r['eta_truth'])
 r['flag']=bool((r['high_absolute'] or r['high_relative']) and (r['location_zero'] or r['scale_large']))
write(O/'flagged_records.csv',[r for r in allrows if r['flag']])
write(O/'all_records.csv',allrows)
(O/'samples.json').write_text(json.dumps(samples,ensure_ascii=False),encoding='utf8')
(O/'r00_inventory.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf8')
(O/'workbook_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf8')
counts=[]
for key in sorted(set((r['project'],r['combination'],r['n'],r['method']) for r in allrows)):
 rs=[r for r in allrows if (r['project'],r['combination'],r['n'],r['method'])==key];f=[r for r in rs if r['flag']]
 counts.append(dict(zip(['project','combination','n','method'],key),lre_version=rs[0]['lre_version'] if key[3]=='LRE' else '',total=len(rs),excel_success=sum(r['excel_success'] for r in rs),flagged=len(f),absolute_flags=sum(r['high_absolute'] and (r['location_zero'] or r['scale_large']) for r in rs),relative_flags=sum(r['high_relative'] and (r['location_zero'] or r['scale_large']) for r in rs),flagged_converged=sum(r['converged'] for r in f),flagged_excel_success=sum(r['excel_success'] for r in f)))
write(O/'counts.csv',counts)
print('TOTAL records',len(allrows),'samples',len(samples),'flags',sum(r['flag'] for r in allrows),flush=True)
