"""Independent read-only workbook-to-record verification; no estimation calls."""
import hashlib,json,math
from pathlib import Path
import openpyxl
HERE=Path(__file__).resolve().parent
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
    cfg=json.loads((HERE/'config.json').read_text(encoding='utf-8'));batch=HERE.parent
    matrix_path=HERE/'中间数据/matrix.json';matrix=json.loads(matrix_path.read_text(encoding='utf-8'))
    original=json.loads((HERE/'输入原件/results.json').read_text(encoding='utf-8'))
    sample_index={(r['n'],r['id']):r['values'] for r in matrix['samples']}
    assert matrix['samples']==original['samples']
    old={(r['n'],r['id'],r['method_id'],r['delta']):r for r in original['results']}
    rows={(r['n'],r['id'],r['method_id'],r['delta']):r for r in matrix['results']}
    for key,r in old.items():assert rows[key]['original_record']==r
    out=HERE.parent/'结果' if HERE.name=='程序' else HERE.parent.parent/'结果'/HERE.name
    book=out/cfg['workbook'];wb=openpyxl.load_workbook(book,read_only=True,data_only=True)
    expected=[f'{prefix}_n{n}' for n in [7,15,30] for prefix in ['生成样本','估计结果']]
    assert wb.sheetnames==expected;records=[];sample_cells=parameter_cells=0;boundary=[]
    for n in [7,15,30]:
        a=list(wb[f'生成样本_n{n}'].iter_rows(values_only=True));b=list(wb[f'估计结果_n{n}'].iter_rows(values_only=True))
        assert len(a)==51 and len(b)==52
        assert [r[0] for r in a[1:]]==[r[0] for r in b[2:]]==list(range(1,51))
        assert list(b[1][1:])==['α','β','γ']*8
        for sid in range(1,51):
            assert list(a[sid][1:])==sample_index[n,sid];sample_cells+=n
            for m in cfg['methods']:
                positions=[i for i,v in enumerate(b[0]) if v==m['excel_label']];assert len(positions)==1
                pos=positions[0];r=rows[n,sid,m['method_id'],m['delta']]
                values=list(b[sid+1][pos:pos+3])
                valid=all(isinstance(v,(int,float)) and math.isfinite(v) for v in values)
                assert valid==r['displayed_success']
                assert values==([r[k] for k in ['beta_hat','eta_hat','gamma_hat']] if valid else ['—']*3)
                parameter_cells+=3
                records.append(dict(n=n,id=sid,method=m['label'],success=valid,values=values if valid else None))
                if valid and r['status']!='success':boundary.append(dict(n=n,id=sid,method=m['label'],status=r['status']))
    wb.close();assert len(records)==1200 and sample_cells==2600 and parameter_cells==3600
    source=dict(comparison_passed=True,errors=[],workbook_sha256=sha(book),matrix_sha256=sha(matrix_path),
      sample_cells_checked=sample_cells,parameter_cells_checked=parameter_cells,reused_records_exact=750,
      sheet_names=expected,method_blocks=[m['excel_label'] for m in cfg['methods']],boundary_records=boundary)
    (HERE/'Excel提取.json').write_text(json.dumps(dict(source=source,records=records),ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (HERE/'复核/Excel一致性.json').write_text(json.dumps(source,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print('EXCEL_CHECK_PASS',cfg['combination'],'6 sheets / 8 blocks / 2600 sample cells / 3600 parameter cells / 750 exact reused records')
if __name__=='__main__':main()
