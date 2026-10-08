"""Extract saved Excel estimates and check them against the frozen matrix. No fitting."""
import hashlib, json, math
from pathlib import Path
import openpyxl

HERE = Path(__file__).resolve().parent

def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def finite(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)

def main():
    cfg = json.loads((HERE/'config.json').read_text(encoding='utf-8'))
    batch = HERE.parent.parent
    workbook = batch/cfg['workbook']; matrix_path = batch/cfg['crosscheck_matrix']
    matrix = json.loads(matrix_path.read_text(encoding='utf-8'))
    indexed = {(r['n'],r['id'],r['method_id'],r['delta']):r for r in matrix['results']}
    sample_index = {(s['n'],s['id']):s['values'] for s in matrix['samples']}
    assert len(indexed)==1200 and len(sample_index)==150
    approved = {tuple(r) for r in cfg['acknowledged_boundary_keys']}
    errors=[]; acknowledged=[]; records=[]; sample_cells=0; parameter_cells=0
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    for n in cfg['sample_sizes']:
        rows=list(wb[f"{cfg['result_sheet_prefix']}_n{n}"].iter_rows(values_only=True))
        assert len(rows)==52 and rows[0][0]=='样本'
        assert list(rows[1][1:])==['α','β','γ']*8
        assert [r[0] for r in rows[2:]]==list(range(1,51))
        headers=[rows[0][1+3*i] for i in range(8)]
        assert len(set(headers))==8 and set(headers)=={m['excel_label'] for m in cfg['methods']}
        for row in rows[2:]:
            id_=int(row[0])
            for method in cfg['methods']:
                i=headers.index(method['excel_label']); values=list(row[1+3*i:4+3*i])
                success=all(finite(v) for v in values)
                failure=all(v in (None,'—','–','-') for v in values)
                key=(n,id_,method['method_id'],method['delta']); ref=indexed[key]
                if not success and not failure:
                    errors.append(dict(key=key,issue='unknown or partial Excel triplet',values=values))
                if success != bool(ref['converged']):
                    errors.append(dict(key=key,issue='Excel display differs from converged'))
                if success != (ref['status']=='success'):
                    x1=min(sample_index[(n,id_)])
                    tolerance=1e-10*max(abs(x1),cfg['truth'][1],1)
                    is_boundary=(key in approved and success and ref['converged'] and
                                 ref['status']=='failure' and 0<=x1-values[2]<=tolerance)
                    diff=dict(key=key,excel_success=success,matrix_converged=ref['converged'],
                              matrix_status=ref['status'],gamma_hat=ref['gamma_hat'],
                              sample_min=x1,gap=x1-ref['gamma_hat'],boundary_tolerance=tolerance)
                    if is_boundary: acknowledged.append(diff)
                    else: errors.append(dict(issue='unacknowledged matrix status difference',**diff))
                if success:
                    for field,value in zip(('beta_hat','eta_hat','gamma_hat'),values):
                        parameter_cells+=1
                        if not finite(ref[field]) or float(value)!=float(ref[field]):
                            errors.append(dict(key=key,issue='numeric mismatch',field=field))
                records.append(dict(n=n,id=id_,method=method['label'],success=success,
                                    values=values if success else None))
        sample_rows=list(wb[f"{cfg['sample_sheet_prefix']}_n{n}"].iter_rows(values_only=True))
        assert len(sample_rows)==51 and [r[0] for r in sample_rows[1:]]==list(range(1,51))
        for row in sample_rows[1:]:
            reference=sample_index[(n,int(row[0]))]; values=list(row[1:n+1])
            assert len(reference)==len(values)==n
            for i,(value,expected) in enumerate(zip(values,reference),1):
                sample_cells+=1
                if not finite(value) or not finite(expected) or float(value)!=float(expected):
                    errors.append(dict(key=[n,row[0],i],issue='sample numeric mismatch'))
    wb.close()
    if {tuple(r['key']) for r in acknowledged} != approved:
        errors.append(dict(issue='acknowledged boundary inventory changed'))
    assert len(records)==1200 and sample_cells==2600
    report=dict(combination=cfg['combination'],workbook=str(workbook),workbook_sha256=sha(workbook),
                crosscheck_matrix=str(matrix_path),matrix_sha256=sha(matrix_path),
                sample_values_compared=sample_cells,successful_parameter_cells_compared=parameter_cells,
                method_records=len(records),success=sum(r['success'] for r in records),
                success_rule=cfg['success_rule'],numeric_comparison='exact float equality',
                acknowledged_status_differences=acknowledged,errors=errors,comparison_passed=not errors)
    (HERE/'数据一致性核对.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    if errors: raise RuntimeError(f'{len(errors)} unexpected differences; no plots produced')
    (HERE/'Excel提取.json').write_text(json.dumps(dict(source=report,records=records),
                                                ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(cfg['combination'], 'EXCEL_CHECK_PASS',sample_cells,parameter_cells,
          'boundary_records',len(acknowledged),flush=True)

if __name__=='__main__': main()
