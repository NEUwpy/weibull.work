"""Read-only audit of case scope, regenerated samples and saved estimate cells.

Workbooks are never saved. JSON report is the only output. --recalculate additionally
checks the first sample of each size using the batch-local dependency snapshot.
"""
from pathlib import Path
import ast
import csv
import hashlib
import importlib.util
import json
import re
import sys

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[1]
sys.path.insert(0, str(REPO / 'python'))
from studies.common.sample import generate_sample


def load_json(path):
    return json.loads(path.read_text(encoding='utf-8-sig'))


def constants(path):
    out = {}
    for node in ast.parse(path.read_text(encoding='utf-8-sig')).body:
        if isinstance(node, ast.Assign) and len(node.targets) == 1 and isinstance(node.targets[0], ast.Name):
            try:
                out[node.targets[0].id] = ast.literal_eval(node.value)
            except (ValueError, TypeError):
                pass
    return out


def saved_estimates(batch):
    data = batch / '结果' / '中间数据'
    table = {}
    payload = data / 'payload.json'
    if payload.exists():
        candidates = [c for c in load_json(payload)['cases'] if c['label'] == batch.parent.name]
        assert len(candidates) == 1
        case = candidates[0]
        for n, blocks in case['results'].items():
            for delta, rows in blocks.items():
                for r in rows:
                    for method in ['MDM','LS','LRE','WMLM','MLM']:
                        v = r[method]
                        table[int(n),r['sample_id'],float(delta),method] = ([v[k] for k in ('shape_hat','scale_hat','location_hat')] if v['converged'] else None)
        return table
    for p in data.glob('**/results.json'):
        for r in load_json(p)['results']:
            method = {'mdm':'MDM','lse':'LS','lre':'LRE','wmle':'WMLM','mle':'MLM'}[r['method_id']]
            for delta in ([r['delta']] if method == 'MDM' else [.1,.15,.2]):
                table[r['n'],r['id'],delta,method] = ([r[k] for k in ('beta_hat','eta_hat','gamma_hat')] if r['converged'] else None)
    for name in ('mdm_estimates.csv','parameter_estimates.csv','other_method_estimates.csv'):
        p = data / name
        if not p.exists():
            continue
        for r in csv.DictReader(p.open(encoding='utf-8-sig',newline='')):
            method = r.get('display_name') or 'MDM'
            for delta in ([float(r['offset'])] if method == 'MDM' else [.1,.15,.2]):
                table[int(r['sample_size']),int(r['sample_id']),delta,method] = ([float(r[k]) for k in ('beta_hat','eta_hat','gamma_hat')] if r['converged'].lower()=='true' else None)
    return table


def main():
    report = {'scope':'参数、样本逐值及保存的结果到Excel逐值核对；不等于历史方法全量独立复算', 'batches':[]}
    for combo in sorted(ROOT.glob('W(*)')):
        truth = tuple(map(float,re.findall(r'[\d.]+',combo.name)))
        for batch in sorted(combo.iterdir()):
            if not batch.is_dir() or '加2000' in batch.name:
                continue
            candidates=[batch/'程序'/n for n in ('generate_cases.py','run_w5_cases.py','run_task.py','run_gamma500.py','fresh_batch.py','generate_samples.py') if (batch/'程序'/n).exists()]
            assert len(candidates)==1,(batch,candidates)
            script=candidates[0]; c=constants(script)
            shape=c.get('SHAPE',c.get('BETA',truth[0]));scale=c.get('SCALE',c.get('ETA',truth[1]))
            locations=c.get('LOCATIONS') or tuple(x['gamma'] for x in c.get('CASES',[])) or (c.get('LOCATION',c.get('GAMMA',truth[2])),)
            assert shape==truth[0] and scale==truth[1] and locations==(truth[2],),(batch,shape,scale,locations)
            seed=c.get('SEED_NAMESPACE',c.get('seed'))
            assert seed is not None,script
            sampler_spec=importlib.util.spec_from_file_location('batch_sampler',batch/'程序/依赖快照/python/studies/common/sample.py')
            sampler=importlib.util.module_from_spec(sampler_spec)
            sampler_spec.loader.exec_module(sampler)
            expected=saved_estimates(batch)
            entry={'batch':batch.relative_to(ROOT).as_posix(),'script':script.name,'truth':truth,'seed':seed,'sample_groups':0,'sample_values':0,'estimate_triples':0,'sample_max_abs_error':0.,'estimate_max_abs_error':0.,'mismatches':[],'spot_recalculations':[]}
            sample_map={}
            for p in sorted((batch/'结果').glob('*.xlsx')):
                before=hashlib.sha256(p.read_bytes()).hexdigest()
                book=load_workbook(p,read_only=True,data_only=True)
                for sheet in book:
                    match=re.search(r'n(\d+)',sheet.title)
                    if not match and not ('SAMPLE_SIZE' in c and '样本' in sheet.title):
                        continue
                    n=int(match[1]) if match else c['SAMPLE_SIZE']; rows=list(sheet.values)
                    if '样本' in sheet.title and '结果' not in sheet.title:
                        for row in rows:
                            if not row or not isinstance(row[0],(int,float)) or not 1<=row[0]<=200:
                                continue
                            sid=int(row[0]);values=list(row[2:2+n]) if 'SAMPLE_SIZE' in c else [v for v in row[1:] if isinstance(v,(int,float))]
                            if len(values)==n+1: values=values[1:]
                            assert len(values)==n,(p,sheet.title,sid,len(values))
                            regenerated=sampler.generate_sample(shape,scale,truth[2],n,sid-1,seed=seed)
                            error=float(np.max(np.abs(np.array(values)-regenerated)))
                            entry['sample_max_abs_error']=max(error,entry['sample_max_abs_error'])
                            entry['sample_groups']+=1;entry['sample_values']+=n
                            if error>1e-9:entry['mismatches'].append({'kind':'sample','n':n,'id':sid,'error':error})
                            sample_map[n,sid]=np.array(values)
                    elif '结果' in sheet.title:
                        delta=None
                        for row in rows:
                            if row and isinstance(row[0],str):
                                d=re.search(r'[δΔ]\s*=\s*(0\.\d+)',row[0])
                                if d:delta=float(d[1])
                            if not row or not isinstance(row[0],(int,float)) or not 1<=row[0]<=50 or delta is None:
                                continue
                            sid=int(row[0])
                            for j,method in enumerate(['MDM','LS','LRE','WMLM','MLM']):
                                key=(n,sid,delta,method); values=row[1+3*j:4+3*j]
                                if key not in expected:
                                    entry['mismatches'].append({'kind':'missing_source','key':key});continue
                                wanted=expected[key];entry['estimate_triples']+=1
                                if wanted is None:
                                    if any(isinstance(v,(int,float)) for v in values):entry['mismatches'].append({'kind':'failure_marker','key':key})
                                elif len(values)!=3 or not all(isinstance(v,(int,float)) for v in values):
                                    entry['mismatches'].append({'kind':'missing_value','key':key})
                                else:
                                    error=float(np.max(np.abs(np.array(values)-wanted)))
                                    entry['estimate_max_abs_error']=max(error,entry['estimate_max_abs_error'])
                                    if not np.allclose(values,wanted,rtol=1e-10,atol=1e-8):entry['mismatches'].append({'kind':'estimate','key':key,'error':error})
                book.close()
                assert hashlib.sha256(p.read_bytes()).hexdigest()==before,p
            if '--recalculate' in sys.argv and expected:
                for key in list(sys.modules):
                    if key == 'base' or key == 'methods' or key.startswith('methods.') or key == 'studies' or key.startswith('studies.'):
                        del sys.modules[key]
                sys.path.insert(0, str(batch/'程序'/'依赖快照'/'python'))
                from studies.common.runner import run_method
                for n in sorted({n for n,sid in sample_map}):
                    for method,display in [('mdm','MDM'),('lse','LS'),('lre','LRE'),('wmle','WMLM'),('mle','MLM')]:
                        raw=run_method(method,sample_map[n,1],**({'offset':.1,'gamma_steps':240} if method=='mdm' else {}))
                        wanted=expected[n,1,.1,display]
                        actual=[raw.get(k) for k in ('beta_hat','eta_hat','gamma_hat')] if raw['converged'] else None
                        same=(wanted is None and actual is None) or (wanted is not None and actual is not None and np.allclose(wanted,actual,rtol=1e-6,atol=1e-5))
                        entry['spot_recalculations'].append({'n':n,'id':1,'method':display,'saved':wanted,'current':actual,'matches':bool(same)})
            wanted_count=200 if 'SAMPLE_SIZE' in c else (150 if batch.name in ('20260921','20260929-新抽样') else 100)
            assert entry['sample_groups']==wanted_count,(batch,entry['sample_groups'],wanted_count)
            report['batches'].append(entry)
            print(json.dumps({k:v for k,v in entry.items() if k not in ('spot_recalculations','mismatches')},ensure_ascii=False),flush=True)
            if entry['mismatches']:print('MISMATCH',json.dumps(entry['mismatches'][:3],ensure_ascii=False),flush=True)
    batch=ROOT/'W(5,1000,500)/20260929-原样本加2000检验'
    data=load_json(batch/'结果/中间数据/translation_data.json')
    path=batch/'结果/原样本加2000_估计结果对照.xlsx'
    book=load_workbook(path,read_only=True,data_only=False)
    for n in (7,15):
        sheet=book[f'估计结果_n{n}']
        for i,r in enumerate([r for r in data['rows'] if r['n']==n],7):
            assert [sheet.cell(i,j).value for j in range(1,4)]==[r['id'],r['method'],r['delta']]
            for j,wanted in enumerate((r['before'] or [None]*3)+(r['after'] or [None]*3),4):
                actual=sheet.cell(i,j).value
                assert (actual is None and wanted is None) or (actual is not None and wanted is not None and np.isclose(actual,wanted,rtol=1e-10,atol=1e-8)),(n,i,j)
    sheet=book['样本对照']
    for i,r in enumerate(data['samples'],7):
        assert sheet.cell(i,4).value==float(r['value'])
        assert sheet.cell(i,5).value==f'=D{i}+$B$3'
        assert sheet.cell(i,6).value==f'=E{i}-D{i}'
    assert sheet['B3'].value==data['shift']==2000
    book.close()
    report['translation_check']={'estimate_rows':len(data['rows']),'sample_values':len(data['samples']),'shift':2000,'saved_values_and_sample_formulas_match':True,'independent_full_recalculation':False}
    (ROOT/'程序结果核对.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print('TOTAL',sum(b['sample_groups'] for b in report['batches']),sum(b['estimate_triples'] for b in report['batches']),sum(len(b['mismatches']) for b in report['batches']),flush=True)


if __name__=='__main__':
    main()
